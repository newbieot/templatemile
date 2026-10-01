"""Extract the operator-provided XLSX into the compact, read-only postcode asset."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook


def build(source, target):
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb['full']
    rows = ws.iter_rows(values_only=True)
    if tuple(next(rows)) != ('postal_code', 'KELURAHAN', 'KECAMATAN', 'KOTA', 'PROVINSI'):
        raise ValueError('Unexpected source columns; inspect the workbook before rebuilding.')
    dictionaries = {key: [] for key in ('districts', 'cities', 'provinces')}
    indexes = {key: {} for key in dictionaries}

    def intern(key, value):
        if value not in indexes[key]:
            indexes[key][value] = len(dictionaries[key])
            dictionaries[key].append(value)
        return indexes[key][value]

    records = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None for value in values):
            continue
        raw, village, district, city, province = values
        postcode = str(int(raw)) if isinstance(raw, (int, float)) and float(raw).is_integer() else str(raw or '').strip()
        if len(postcode) != 5 or not postcode.isdigit():
            raise ValueError(f'Invalid postcode at source row {row_number}')
        labels = [str(value or '').strip() for value in (village, district, city, province)]
        if not all(labels):
            raise ValueError(f'Missing geography at source row {row_number}')
        records.append([postcode, labels[0], intern('districts', labels[1]), intern('cities', labels[2]), intern('provinces', labels[3])])
    wb.close()
    payload = {
        'formatVersion': 1,
        'source': {'file': source.name, 'sheet': 'full', 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()},
        'recordCount': len(records),
        'dictionaries': dictionaries,
        'rows': records,
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    print(json.dumps({'records': len(records), 'bytes': target.stat().st_size, 'output': str(target)}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('target', type=Path)
    args = parser.parse_args()
    build(args.source, args.target)
